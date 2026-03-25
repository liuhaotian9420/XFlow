"""
Integration tests: sql-generator output contract → dataworks single-file execution path.

The sql-generator skill is expected to produce one ``.sql`` file that dataworks
(``run_sql_export.py``) can consume as ``file_path``. We never call a live LLM here.

What is validated (without real ODPS):
- Each benchmark ``truth_path`` from ``sql_generator_eval_cases.json`` exists and is UTF-8.
- The same files flow through dataworks ``read_sql_file`` → mocked ``run_odps_sql`` → Excel export.
- Optional DuckDB cache path still runs after Excel (one case).

Run::
    uv run pytest tests/test_integration_sql_generator_dataworks.py -v
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import pandas as pd
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
CASES_JSON = REPO_ROOT / "tests" / "sql_generator_eval_cases.json"


def _load_case_ids_and_paths() -> list[tuple[str, str]]:
    raw = json.loads(CASES_JSON.read_text(encoding="utf-8"))
    out: list[tuple[str, str]] = []
    for row in raw:
        cid = row.get("id")
        tp = row.get("truth_path")
        if isinstance(cid, str) and isinstance(tp, str):
            out.append((cid, tp))
    return out


def _import_dataworks_module() -> Any:
    import importlib.util

    path = REPO_ROOT / ".agents" / "skills" / "dataworks" / "scripts" / "run_sql_export.py"
    spec = importlib.util.spec_from_file_location("dataworks_run_sql_export_sg", str(path))
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Failed to load dataworks from {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _mock_odps_frame() -> pd.DataFrame:
    """Deterministic result set standing in for ODPS tunnel output."""
    return pd.DataFrame(
        {
            "loan_date": ["2025-01-01", "2025-01-02"],
            "flag": ["自营总计", "新客总体"],
            "day_loan_amt": [100.0, 200.5],
            "day_loan_cnt": [10, 20],
        }
    )


def _run_dataworks_main(argv: list[str], mod: Any) -> None:
    old = sys.argv[:]
    sys.argv = argv
    try:
        mod.main()
    finally:
        sys.argv = old


def _verify_excel_columns(xlsx_path: Path, expected: list[str]) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    headers_set = {h for h in headers if isinstance(h, str)}
    missing = [c for c in expected if c not in headers_set]
    if missing:
        raise AssertionError(f"Excel missing columns {missing}; got {headers}")


@pytest.mark.integration
@pytest.mark.parametrize("case_id,truth_rel", _load_case_ids_and_paths())
def test_truth_sql_flows_to_dataworks_excel(
    case_id: str,
    truth_rel: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Ground-truth SQL (proxy for sql-generator output) is accepted by dataworks and yields Excel."""
    truth_path = (REPO_ROOT / truth_rel).resolve()
    assert truth_path.is_file(), f"Missing truth SQL: {truth_path}"
    sql_text = truth_path.read_text(encoding="utf-8")
    assert sql_text.strip(), f"Empty truth file: {truth_path}"

    mod = _import_dataworks_module()
    mock_df = _mock_odps_frame()
    monkeypatch.setattr(
        mod,
        "resolve_odps_config",
        lambda _args: ("dummy_id", "dummy_key", "dummy_project", "dummy_endpoint"),
        raising=True,
    )
    monkeypatch.setattr(
        mod,
        "run_odps_sql",
        lambda *_a, **_k: mock_df.copy(),
        raising=True,
    )

    staged_sql = tmp_path / f"{case_id}_from_generator.sql"
    staged_sql.write_text(sql_text, encoding="utf-8")

    excel_dir = tmp_path / "excel"
    _run_dataworks_main(
        [
            "run_sql_export.py",
            str(staged_sql),
            "--save_path",
            str(excel_dir),
        ],
        mod,
    )

    xlsx = excel_dir / f"{staged_sql.stem}.xlsx"
    assert xlsx.is_file(), f"Expected Excel at {xlsx}"
    _verify_excel_columns(xlsx, expected=list(mock_df.columns))


@pytest.mark.integration
def test_sql_generator_to_dataworks_duckdb_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """After Excel, optional DuckDB write path works with a sql-generator-shaped file."""
    mod = _import_dataworks_module()
    mock_df = _mock_odps_frame()
    monkeypatch.setattr(
        mod,
        "resolve_odps_config",
        lambda _args: ("dummy_id", "dummy_key", "dummy_project", "dummy_endpoint"),
        raising=True,
    )
    monkeypatch.setattr(
        mod,
        "run_odps_sql",
        lambda *_a, **_k: mock_df.copy(),
        raising=True,
    )

    # Smallest truth artifact for a real on-disk SQL text
    truth = REPO_ROOT / ".agents/skills/sql-generator/references/sql代码/dws_inloan_loan_risk_stat_df.sql"
    assert truth.is_file()
    staged = tmp_path / "risk_stat_handoff.sql"
    staged.write_text(truth.read_text(encoding="utf-8"), encoding="utf-8")

    excel_dir = tmp_path / "excel"
    db_path = tmp_path / "cache.duckdb"
    table = "sg_dw_handoff"

    _run_dataworks_main(
        [
            "run_sql_export.py",
            str(staged),
            "--save_path",
            str(excel_dir),
            "--duckdb-path",
            str(db_path),
            "--duckdb-table",
            table,
            "--lifecycle",
            "7",
        ],
        mod,
    )

    xlsx = excel_dir / f"{staged.stem}.xlsx"
    assert xlsx.is_file()

    import duckdb  # type: ignore

    con = duckdb.connect(str(db_path))
    try:
        tables = [r[0] for r in con.execute("SHOW TABLES").fetchall()]
        assert table in tables
        n = con.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
        assert int(n) == len(mock_df)
    finally:
        con.close()


@pytest.mark.integration
def test_skill_cross_reference_paths_exist() -> None:
    """Lightweight contract: both SKILL files exist and sql-generator names dataworks."""
    sg = (REPO_ROOT / ".agents/skills/sql-generator/SKILL.md").read_text(encoding="utf-8")
    dw = (REPO_ROOT / ".agents/skills/dataworks/SKILL.md").read_text(encoding="utf-8")
    assert "dataworks" in sg.lower()
    assert "sql" in dw.lower()
    assert (REPO_ROOT / ".agents/skills/dataworks/scripts/run_sql_export.py").is_file()
