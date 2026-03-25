"""
Integration test suite for the interaction between:
1) @.agents/skills/scorecardpy
2) @.agents/skills/dataworks (run_sql_export.py)

This repo does not use pytest; we provide a runnable script that returns
exit code 0 on success, non-zero on failure.

Design goals (per your request):
1) Excel path is tested first.
2) DuckDB output is tested after Excel passes.

Because real ODPS execution is not feasible in CI/local unit tests,
we monkeypatch dataworks' ODPS execution to return a deterministic DataFrame.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parent.parent


def _run_scorecardpy_to_get_x_cols(tmp_path: Path) -> tuple[Path, list[str]]:
    """
    Run scorecardpy on DuckDB germancredit() sample and return:
    - out_dir
    - x_cols (selected by the skill)
    """
    import duckdb  # type: ignore
    import scorecardpy as sc  # type: ignore

    run_script = REPO_ROOT / ".agents" / "skills" / "scorecardpy" / "scripts" / "run_scorecardpy.py"

    # Build a minimal dev sample in a temp DuckDB.
    db_path = tmp_path / "scorecardpy_smoke.duckdb"
    out_dir = tmp_path / "scorecardpy_out"

    raw = sc.germancredit()
    df = raw.drop(columns=["creditability"]).copy()
    df.insert(0, "entity_id", range(len(df)))
    df["target"] = (raw["creditability"] == "bad").astype(int)

    con = duckdb.connect(str(db_path))
    try:
        con.register("df_in", df)
        con.execute("CREATE TABLE dev_smoke_test AS SELECT * FROM df_in")
    finally:
        con.close()

    cmd = [
        sys.executable,
        str(run_script),
        "--db-path",
        str(db_path),
        "--table",
        "dev_smoke_test",
        "--target",
        "target",
        "--id-col",
        "entity_id",
        "--out-dir",
        str(out_dir),
    ]
    proc = subprocess.run(cmd, cwd=str(REPO_ROOT), capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise RuntimeError(
            "scorecardpy run_scorecardpy.py failed:\n"
            f"STDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}"
        )

    run_summary_path = out_dir / "run_summary.json"
    if not run_summary_path.is_file():
        raise RuntimeError(f"Missing scorecardpy artifact: {run_summary_path}")

    summary: dict[str, Any] = json.loads(run_summary_path.read_text(encoding="utf-8"))
    x_cols = summary.get("x_cols")
    if not isinstance(x_cols, list) or not all(isinstance(x, str) for x in x_cols):
        raise RuntimeError(f"Invalid x_cols in {run_summary_path}: {x_cols!r}")
    if not x_cols:
        raise RuntimeError("scorecardpy returned empty x_cols")

    return out_dir, x_cols


def _import_dataworks_module():
    """
    Import dataworks' run_sql_export.py via file path so we can monkeypatch it.
    """
    import importlib.util

    path = REPO_ROOT / ".agents" / "skills" / "dataworks" / "scripts" / "run_sql_export.py"
    spec = importlib.util.spec_from_file_location("dataworks_run_sql_export", str(path))
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Failed to import dataworks module from {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _verify_excel_headers(xlsx_path: Path, expected_cols: list[str]) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = [ws.cell(1, col_idx).value for col_idx in range(1, ws.max_column + 1)]
    headers_set = {h for h in headers if isinstance(h, str)}

    missing = [c for c in expected_cols if c not in headers_set]
    if missing:
        raise AssertionError(f"Excel missing expected columns: {missing}. Headers={headers}")


def _verify_duckdb_table_and_lifecycle(
    db_path: Path,
    table_name: str,
    lifecycle_days: int,
    old_table_name: str,
) -> None:
    import duckdb  # type: ignore

    con = duckdb.connect(str(db_path))
    try:
        tables = [r[0] for r in con.execute("SHOW TABLES").fetchall()]
        if table_name not in tables:
            raise AssertionError(f"DuckDB table not found: {table_name}")

        row = con.execute(
            "SELECT table_name, lifecycle_days FROM __table_lifecycle WHERE table_name = ?",
            [table_name],
        ).fetchone()
        if row is None:
            raise AssertionError(f"Lifecycle metadata missing for: {table_name}")
        if row[1] != lifecycle_days:
            raise AssertionError(f"lifecycle_days mismatch: got {row[1]}, expected {lifecycle_days}")

        # Contract: expired cached tables should be dropped.
        old_exists = con.execute(
            "SELECT 1 FROM information_schema.tables WHERE table_name = ? LIMIT 1",
            [old_table_name],
        ).fetchone()
        if old_exists is not None:
            raise AssertionError(f"Expired old table still exists: {old_table_name}")
    finally:
        con.close()


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)

        # 1) Run scorecardpy and obtain handpick candidate columns.
        _, x_cols = _run_scorecardpy_to_get_x_cols(tmp_path)

        # 2) Import dataworks and monkeypatch ODPS execution.
        dataworks = _import_dataworks_module()

        # Make ODPS config resolution deterministic + avoid env/credentials.
        dataworks.resolve_odps_config = lambda args: ("dummy_access_id", "dummy_access_key", "dummy_project", "dummy_endpoint")

        # Mock ODPS execution to return a deterministic DataFrame.
        # Include entity_id so downstream "scoring-like" results feel realistic.
        import pandas as pd  # type: ignore

        df = pd.DataFrame(
            {
                "entity_id": [1, 2, 3],
                **{c: [0.1, 0.2, 0.3] for c in x_cols[:10]},  # keep it wide but bounded
                "score": [0.9, 0.1, 0.5],
            }
        )

        # Ensure returned columns include at least the handpicked subset.
        expected_cols = list(df.columns)

        dataworks.run_odps_sql = lambda *args, **kwargs: df

        # Prepare a dummy SQL file for dataworks to read.
        sql_path = tmp_path / "dummy_odps_query.sql"
        sql_path.write_text("SELECT 1 AS ok;")

        # -----------------------
        # Excel-first phase
        # -----------------------
        excel_out_dir = tmp_path / "excel_out"
        xlsx_name = "dummy_odps_query.xlsx"
        dataworks_sql_argv = [
            str(sql_path),
            "--save_path",
            str(excel_out_dir),
        ]

        sys.argv = ["run_sql_export.py", *dataworks_sql_argv]
        proc_out: str = ""
        try:
            dataworks.main()
        except SystemExit as e:
            return int(e.code or 1)

        xlsx_path = excel_out_dir / xlsx_name
        if not xlsx_path.is_file():
            raise AssertionError(f"Excel output missing: {xlsx_path}")
        _verify_excel_headers(xlsx_path, expected_cols=expected_cols)

        # -----------------------
        # DuckDB phase
        # -----------------------
        db_path = tmp_path / "integration_cache.duckdb"
        duckdb_table = "tmp_integration_result"
        old_table = "old_expired_cache"
        lifecycle_days = 2

        # Pre-seed an expired cache entry + table to verify cleanup semantics.
        import duckdb  # type: ignore
        from datetime import datetime, timedelta

        con = duckdb.connect(str(db_path))
        try:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS __table_lifecycle (
                    table_name VARCHAR PRIMARY KEY,
                    created_at TIMESTAMP,
                    expires_at TIMESTAMP,
                    lifecycle_days INTEGER
                )
                """
            )
            created_at = datetime.now() - timedelta(days=5)
            expires_at = datetime.now() - timedelta(days=1)
            con.execute(
                "INSERT OR REPLACE INTO __table_lifecycle(table_name, created_at, expires_at, lifecycle_days) VALUES (?, ?, ?, ?)",
                [old_table, created_at, expires_at, 1],
            )
            con.execute(f"CREATE OR REPLACE TABLE {old_table} AS SELECT 1 AS a")
        finally:
            con.close()

        # Re-run dataworks to also write DuckDB cache.
        duckdb_out_dir = tmp_path / "excel_out2"
        sys.argv = [
            "run_sql_export.py",
            str(sql_path),
            "--save_path",
            str(duckdb_out_dir),
            "--duckdb-path",
            str(db_path),
            "--duckdb-table",
            duckdb_table,
            "--lifecycle",
            str(lifecycle_days),
        ]
        try:
            dataworks.main()
        except SystemExit as e:
            return int(e.code or 1)

        _verify_duckdb_table_and_lifecycle(
            db_path=db_path,
            table_name=duckdb_table,
            lifecycle_days=lifecycle_days,
            old_table_name=old_table,
        )

        print("INTEGRATION_SCORECARDBPY_DATAWORKS_OK")
        return 0


if __name__ == "__main__":
    raise SystemExit(main())

