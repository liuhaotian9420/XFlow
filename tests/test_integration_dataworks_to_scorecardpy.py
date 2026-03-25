"""
Integration test suite for the interaction between:
1) @.agents/skills/dataworks (run_sql_export.py -> Excel -> DuckDB optional cache)
2) @.agents/skills/scorecardpy (run_scorecardpy.py -> artifacts)

This test follows the corrected联调顺序:
dataworks 出数据（优先 Excel） -> 再交由 scorecardpy 来分析

Notes:
- Real ODPS execution is not feasible here, so we monkeypatch `dataworks.run_odps_sql`
  to return a deterministic pandas DataFrame.
- We run dataworks twice to respect "Excel first, then DuckDB出表":
  1) Excel-only run (no --duckdb-table) and verify .xlsx headers.
  2) DuckDB cache run (with --duckdb-table) and verify cache table + lifecycle cleanup.
"""

from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parent.parent


def _make_mock_df(n: int = 400, seed: int = 42) -> "Any":
    import numpy as np
    import pandas as pd

    rng = np.random.default_rng(seed)
    entity_id = np.arange(n, dtype=int)
    target = (rng.random(n) < 0.25).astype(int)

    # Numeric features correlated with target so var_filter / woebin have signal.
    feature_num_1 = (target * 2.0 + rng.normal(0, 1.0, n)).astype(float)
    feature_num_2 = (rng.lognormal(mean=1.0 + target * 0.2, sigma=0.5, size=n)).astype(float)
    feature_num_3 = (rng.normal(loc=0.0, scale=1.0, size=n) + target * -0.8).astype(float)

    # Inject small missingness (keep it low to avoid var_filter dropping everything).
    miss_mask = rng.random(n) < 0.05
    feature_num_3[miss_mask] = np.nan

    # Categorical features with target-dependent probabilities.
    cats1 = np.array(["A", "B", "C"], dtype=object)
    cat1 = np.empty(n, dtype=object)
    for i in range(n):
        if target[i] == 1:
            probs = [0.2, 0.6, 0.2]
        else:
            probs = [0.5, 0.3, 0.2]
        cat1[i] = rng.choice(cats1, p=probs)

    cats2 = np.array(["X", "Y"], dtype=object)
    cat2 = np.empty(n, dtype=object)
    for i in range(n):
        probs = [0.7, 0.3] if target[i] == 1 else [0.3, 0.7]
        cat2[i] = rng.choice(cats2, p=probs)

    df = pd.DataFrame(
        {
            "entity_id": entity_id,
            "target": target,
            "feature_num_1": feature_num_1,
            "feature_num_2": feature_num_2,
            "feature_num_3": feature_num_3,
            "feature_cat_1": cat1,
            "feature_cat_2": cat2,
        }
    )
    return df


def _import_dataworks_module():
    import importlib.util

    path = REPO_ROOT / ".agents" / "skills" / "dataworks" / "scripts" / "run_sql_export.py"
    spec = importlib.util.spec_from_file_location("dataworks_run_sql_export", str(path))
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Failed to import dataworks module from {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _verify_excel_headers(xlsx_path: Path, expected_cols: list[str]) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    # header is first row
    max_col = ws.max_column
    headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]
    headers_set = {h for h in headers if isinstance(h, str)}

    missing = [c for c in expected_cols if c not in headers_set]
    if missing:
        raise AssertionError(f"Excel missing expected columns: {missing}. Headers={headers}")


def _seed_expired_cache(db_path: Path, old_table: str) -> None:
    import duckdb  # type: ignore

    con = duckdb.connect(str(db_path))
    try:
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS __table_lifecycle (
                table_name VARCHAR PRIMARY KEY,
                created_at TIMESTAMP,
                expires_at TIMESTAMP,
                lifecycle_days INTEGER
            )
            """
        )
        created_at = datetime.now() - timedelta(days=5)
        expires_at = datetime.now() - timedelta(days=1)
        con.execute(
            "INSERT OR REPLACE INTO __table_lifecycle(table_name, created_at, expires_at, lifecycle_days) VALUES (?, ?, ?, ?)",
            [old_table, created_at, expires_at, 1],
        )
        con.execute(f'CREATE OR REPLACE TABLE {old_table} AS SELECT 1 AS a')
    finally:
        con.close()


def _verify_duckdb_table_and_lifecycle(
    db_path: Path,
    table_name: str,
    lifecycle_days: int,
    old_table_name: str,
) -> None:
    import duckdb  # type: ignore

    con = duckdb.connect(str(db_path))
    try:
        tables = [r[0] for r in con.execute("SHOW TABLES").fetchall()]
        if table_name not in tables:
            raise AssertionError(f"DuckDB table not found: {table_name}")

        row = con.execute(
            "SELECT table_name, lifecycle_days FROM __table_lifecycle WHERE table_name = ?",
            [table_name],
        ).fetchone()
        if row is None:
            raise AssertionError(f"Lifecycle metadata missing for: {table_name}")
        if row[1] != lifecycle_days:
            raise AssertionError(f"lifecycle_days mismatch: got {row[1]}, expected {lifecycle_days}")

        old_exists = con.execute(
            "SELECT 1 FROM information_schema.tables WHERE table_name = ? LIMIT 1",
            [old_table_name],
        ).fetchone()
        if old_exists is not None:
            raise AssertionError(f"Expired old table still exists: {old_table_name}")
    finally:
        con.close()


def _run_scorecardpy(db_path: Path, table: str, out_dir: Path) -> None:
    run_script = REPO_ROOT / ".agents" / "skills" / "scorecardpy" / "scripts" / "run_scorecardpy.py"
    export_script = REPO_ROOT / ".agents" / "skills" / "scorecardpy" / "scripts" / "export_breaks.py"

    cmd = [
        sys.executable,
        str(run_script),
        "--db-path",
        str(db_path),
        "--table",
        table,
        "--target",
        "target",
        "--id-col",
        "entity_id",
        "--out-dir",
        str(out_dir),
        "--save-plots",
        "--plots-max",
        "30",
    ]
    proc = subprocess.run(cmd, cwd=str(REPO_ROOT), capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise RuntimeError(f"run_scorecardpy.py failed (exit {proc.returncode})\nSTDERR:\n{proc.stderr}\nSTDOUT:\n{proc.stdout}")

    bins_pkl = out_dir / "bins.pkl"
    if not bins_pkl.is_file():
        raise RuntimeError(f"Missing {bins_pkl}")

    plots_dir = out_dir / "plots"
    if not plots_dir.is_dir():
        raise RuntimeError(f"Missing plots dir: {plots_dir}")
    pngs = list(plots_dir.glob("*.png"))
    htmls = list(plots_dir.glob("*.html"))
    if not pngs:
        raise RuntimeError(f"No .png plots produced under {plots_dir}")
    if not htmls:
        raise RuntimeError(f"No .html plots produced under {plots_dir}")
    # Basic self-contained html check: should embed base64 png
    sample_html = htmls[0].read_text(encoding="utf-8")
    if "data:image/png;base64" not in sample_html:
        raise RuntimeError(f"Expected self-contained base64 html, got: {htmls[0]}")

    proc2 = subprocess.run(
        [sys.executable, str(export_script), "--bins-pkl", str(bins_pkl), "--out-dir", str(out_dir)],
        cwd=str(REPO_ROOT),
        capture_output=True,
        text=True,
        check=False,
    )
    if proc2.returncode != 0:
        raise RuntimeError(f"export_breaks.py failed (exit {proc2.returncode})\nSTDERR:\n{proc2.stderr}\nSTDOUT:\n{proc2.stdout}")

    if not (out_dir / "breaks_export.json").is_file():
        raise RuntimeError("Missing breaks_export.json")
    if not (out_dir / "selected_variables.json").is_file():
        raise RuntimeError("Missing selected_variables.json")


def main() -> int:
    df = _make_mock_df()
    expected_cols = list(df.columns)

    dataworks = _import_dataworks_module()
    # Avoid real ODPS credentials.
    dataworks.resolve_odps_config = lambda args: (
        "dummy_access_id",
        "dummy_access_key",
        "dummy_project",
        "dummy_endpoint",
    )
    dataworks.run_odps_sql = lambda *args, **kwargs: df

    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)

        sql_path = tmp_path / "dummy_odps_query.sql"
        sql_path.write_text("SELECT 1 AS ok;")

        # -----------------------
        # Excel-first phase
        # -----------------------
        excel_out_dir = tmp_path / "excel_out"
        sys.argv = [
            "run_sql_export.py",
            str(sql_path),
            "--save_path",
            str(excel_out_dir),
        ]
        try:
            dataworks.main()
        except SystemExit as e:
            return int(e.code or 1)

        xlsx_path = excel_out_dir / "dummy_odps_query.xlsx"
        if not xlsx_path.is_file():
            raise AssertionError(f"Excel output missing: {xlsx_path}")
        _verify_excel_headers(xlsx_path, expected_cols=expected_cols)

        # -----------------------
        # DuckDB phase (after Excel passes)
        # -----------------------
        db_path = tmp_path / "integration_cache.duckdb"
        duckdb_table = "tmp_integration_result"
        old_table = "old_expired_cache"
        lifecycle_days = 2
        _seed_expired_cache(db_path, old_table)

        duckdb_excel_out_dir = tmp_path / "excel_out2"
        sys.argv = [
            "run_sql_export.py",
            str(sql_path),
            "--save_path",
            str(duckdb_excel_out_dir),
            "--duckdb-path",
            str(db_path),
            "--duckdb-table",
            duckdb_table,
            "--lifecycle",
            str(lifecycle_days),
        ]
        try:
            dataworks.main()
        except SystemExit as e:
            return int(e.code or 1)

        _verify_duckdb_table_and_lifecycle(
            db_path=db_path,
            table_name=duckdb_table,
            lifecycle_days=lifecycle_days,
            old_table_name=old_table,
        )

        # -----------------------
        # Scorecardpy phase
        # -----------------------
        score_out_dir = tmp_path / "score_out"
        _run_scorecardpy(db_path=db_path, table=duckdb_table, out_dir=score_out_dir)

        print("INTEGRATION_DATAWORKS_TO_SCORECARDBPY_OK")
        return 0


if __name__ == "__main__":
    raise SystemExit(main())

